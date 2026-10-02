from __future__ import annotations

import io

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

from snoska.domain.documents.structure import Alignment, BlockKind
from snoska.domain.documents.version import ExtractionMethod, ParseFailureCode
from snoska.infrastructure.parsing.docx.numbering import format_number
from snoska.infrastructure.parsing.docx.parser import DocxParser
from snoska.ports.document_parsing import DocumentParseError


def build_docx() -> bytes:
    document = Document()
    section = document.sections[0]
    section.left_margin = Mm(30)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    document.add_heading("Общие положения", level=1)
    body = document.add_paragraph("1.1. Настоящая политика определяет порядок обработки.")
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = Mm(12.5)
    for text in ("первый пункт;", "второй пункт."):
        document.add_paragraph(text, style="List Number")
    title = document.add_paragraph()
    run = title.add_run("ПРАВА СУБЪЕКТА")
    run.bold = True
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Цель"
    table.cell(0, 2).text = "Срок"
    table.cell(1, 0).text = "Доставка"
    table.cell(1, 1).text = "ФИО"
    table.cell(1, 2).text = "3 года"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_structure_formatting_numbering_and_tables() -> None:
    parsed = await DocxParser().parse(build_docx())
    assert parsed.method is ExtractionMethod.DOCX
    kinds = [block.kind for block in parsed.blocks]
    assert kinds == [
        BlockKind.HEADING,
        BlockKind.LIST_ITEM,
        BlockKind.LIST_ITEM,
        BlockKind.LIST_ITEM,
        BlockKind.HEADING,
        BlockKind.TABLE,
    ]
    heading, clause, first, second, caps, table = parsed.blocks
    assert heading.level == 1
    assert (clause.label, clause.text) == (
        "1.1.",
        "Настоящая политика определяет порядок обработки.",
    )
    assert clause.paragraph is not None
    assert clause.paragraph.alignment is Alignment.JUSTIFY
    assert clause.paragraph.first_line_indent_mm == pytest.approx(12.5, abs=0.1)
    assert clause.fonts[0].family == "Times New Roman"
    assert clause.fonts[0].size_pt == 14
    assert (first.label, second.label) == ("1.", "2.")
    assert caps.text == "ПРАВА СУБЪЕКТА"
    assert table.table is not None
    assert table.table.rows[0][0].col_span == 2
    assert parsed.layout.pages[0].margin_left_mm == pytest.approx(30, abs=0.1)


async def test_corrupted_file_is_a_parse_error_not_a_crash() -> None:
    with pytest.raises(DocumentParseError) as error:
        await DocxParser().parse(b"PK\x03\x04 not really a document")
    assert error.value.code is ParseFailureCode.CORRUPTED_FILE


@pytest.mark.parametrize(
    ("value", "fmt", "label"),
    [
        (3, "decimal", "3"),
        (2, "russianLower", "б"),
        (30, "russianLower", "бб"),
        (4, "upperRoman", "IV"),
    ],
)
def test_word_number_formats(value: int, fmt: str, label: str) -> None:
    assert format_number(value, fmt) == label
