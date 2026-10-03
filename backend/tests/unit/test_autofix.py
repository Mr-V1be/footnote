"""Corrections of formatting: their rewrites and the DOCX they produce."""

from __future__ import annotations

import io

import docx
import pytest
from docx.shared import Pt

from snoska.domain.autofix.actions import FixPlan, ReplaceText, SetFontSize
from snoska.domain.autofix.planning import arabic_label, initials_first, standard_date
from snoska.domain.documents.structure import DocumentStructure
from snoska.infrastructure.autofix.docx_patcher import DocxPatcher
from snoska.infrastructure.parsing.docx.parser import DocxParser


@pytest.mark.parametrize(
    ("value", "expected"),
    [("1.2.2026", "01.02.2026"), ("2026-02-03", "03.02.2026"), ("5/6/26", "05.06.2026")],
)
def test_dates_take_the_standard_spelling(value: str, expected: str) -> None:
    assert standard_date(value) == expected


def test_signatures_and_section_numbers_are_rewritten() -> None:
    assert initials_first("Петров И.И.") == "И.И. Петров"
    assert initials_first("И.И. Петров") is None
    assert [arabic_label(label) for label in ("I.", "IV.", "XII", "1.")] == ["1.", "4.", "12", None]


def sample() -> bytes:
    document = docx.Document()
    first = document.add_paragraph()
    for text in ("Утверждено ", "1.2.", "2026 директором Петров И.И."):
        first.add_run(text).font.size = Pt(11)
    document.add_paragraph("Второй абзац основного текста.").runs[0].font.size = Pt(10)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_the_patcher_rewrites_across_runs_and_resizes_the_text() -> None:
    content = sample()
    structure = await DocxParser().parse(content)
    parsed = DocumentStructure(blocks=structure.blocks, layout=structure.layout)
    plan = FixPlan(
        actions=(
            ReplaceText("formatting.date_format", 0, "1.2.2026", "01.02.2026"),
            ReplaceText("formatting.signature_order", 0, "Петров И.И.", "И.И. Петров"),
            SetFontSize("formatting.font_size", (0, 1), 12, 14),
            ReplaceText("formatting.date_format", 1, "3.3.2026", "03.03.2026"),
        ),
        skipped=(),
    )
    result = await DocxPatcher().apply(content, plan, parsed)
    assert [change.check_code for change in result.changes] == [
        "formatting.date_format",
        "formatting.signature_order",
        "formatting.font_size",
    ]
    assert len(result.missed) == 1, "text that is not there is reported, not invented"
    reread = docx.Document(io.BytesIO(result.content))
    assert reread.paragraphs[0].text == "Утверждено 01.02.2026 директором И.И. Петров"
    assert {run.font.size.pt for paragraph in reread.paragraphs for run in paragraph.runs} == {12}
    original = docx.Document(io.BytesIO(content))
    assert original.paragraphs[0].text.startswith("Утверждено 1.2.2026"), "the source is untouched"
