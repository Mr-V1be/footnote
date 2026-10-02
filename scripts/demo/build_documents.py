"""Builds the demo documents and the reference set into backend/demo.

Run with the backend environment (python-docx, pypdfium2 and Pillow are its
dependencies); LibreOffice converts DOCX to PDF:

    backend/.venv/bin/python scripts/demo/build_documents.py [path prefix...]

With path prefixes (``reference/``) only the matching documents are rebuilt.
"""

from __future__ import annotations

import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pypdfium2
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from PIL import Image, ImageFilter

sys.path.insert(0, str(Path(__file__).parent))

import content_polar_other
import content_polar_policy
import content_port
import reference_contracts
import reference_orders
import reference_policies
import reference_policy_full
import reference_short
from model import DemoDocument, PageBreak, PageSetup, Para, Style, Table

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "backend" / "demo"
SCAN_DPI = 200
SCAN_SKEW_DEGREES = 0.6
SCAN_NOISE_SIGMA = 18
SCAN_NOISE_SHARE = 0.05
SCAN_BLUR = 0.4
SCAN_JPEG_QUALITY = 70

DOCUMENTS: tuple[DemoDocument, ...] = (
    content_polar_policy.POLICY,
    content_polar_other.CONSENT,
    content_polar_other.ORDER,
    content_polar_other.CONTRACT,
    content_polar_other.REGULATION_V1,
    content_polar_other.REGULATION_V2,
    content_port.POLICY,
    content_port.EMPLOYEE_CONSENT,
    # The reference set: what the checks are measured against.
    reference_policy_full.POLICY_FULL,
    reference_policies.POLICY_FLAWED,
    reference_policies.POLICY_SCAN,
    reference_policies.REGULATION_FLAWED,
    reference_contracts.CONSENT_FULL,
    reference_contracts.CONSENT_FLAWED,
    reference_contracts.ASSIGNMENT_FULL,
    reference_contracts.ASSIGNMENT_FLAWED,
    reference_orders.ORDER_CORRECT,
    reference_orders.ORDER_FLAWED,
    reference_short.POLICY_SITE,
    reference_short.CONSENT_NEWSLETTER,
    reference_short.ASSIGNMENT_CLAUSE,
    reference_short.ORDER_UNSIGNED,
)

_ALIGN = {
    Style.TITLE: WD_ALIGN_PARAGRAPH.CENTER,
    Style.SUBTITLE: WD_ALIGN_PARAGRAPH.CENTER,
    Style.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Style.HEADING: WD_ALIGN_PARAGRAPH.LEFT,
    Style.BODY: WD_ALIGN_PARAGRAPH.JUSTIFY,
    Style.CLAUSE: WD_ALIGN_PARAGRAPH.JUSTIFY,
    Style.BULLET: WD_ALIGN_PARAGRAPH.JUSTIFY,
    Style.NUMBERED: WD_ALIGN_PARAGRAPH.JUSTIFY,
    Style.RIGHT: WD_ALIGN_PARAGRAPH.LEFT,
    Style.PLAIN: WD_ALIGN_PARAGRAPH.LEFT,
    Style.SIGNATURE: WD_ALIGN_PARAGRAPH.LEFT,
}
_LIST_STYLES = {Style.BULLET: "List Bullet", Style.NUMBERED: "List Number"}
_BOLD = {Style.TITLE, Style.SUBTITLE, Style.HEADING}
_INDENTED = {Style.BODY, Style.CLAUSE, Style.HEADING}
APPROVAL_BLOCK_INDENT_MM = 100


def _write_docx(spec: DemoDocument, target: Path) -> None:
    page: PageSetup = spec.page
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(page.margin_top), Mm(page.margin_bottom)
    section.left_margin, section.right_margin = Mm(page.margin_left), Mm(page.margin_right)
    for style_name in ("Normal", "List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = page.font
        style.font.size = Pt(page.size)
    for element in spec.elements:
        if isinstance(element, PageBreak):
            document.add_page_break()
        elif isinstance(element, Table):
            _write_table(document, element, page)
        else:
            _write_paragraph(document, element, page)
    document.core_properties.author = ""
    document.core_properties.comments = ""
    document.save(str(target))


def _write_paragraph(document: Document, element: Para, page: PageSetup) -> None:
    paragraph = document.add_paragraph(style=_LIST_STYLES.get(element.style))
    run = paragraph.add_run(element.text)
    run.bold = element.style in _BOLD
    fmt = paragraph.paragraph_format
    fmt.alignment = _ALIGN[element.style]
    fmt.line_spacing = page.line_spacing
    fmt.space_before = Pt(12 if element.style is Style.HEADING else 0)
    fmt.space_after = Pt(12 if element.style is Style.SUBTITLE else 0)
    if element.style in _INDENTED:
        fmt.first_line_indent = Mm(page.first_line_indent)
    if element.style is Style.RIGHT:
        fmt.left_indent = Mm(APPROVAL_BLOCK_INDENT_MM)
    if element.style is Style.HEADING:
        fmt.keep_with_next = True


def _write_table(document: Document, element: Table, page: PageSetup) -> None:
    columns = max(len(row) for row in element.rows)
    table = document.add_table(rows=len(element.rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if element.borders:
        table.style = document.styles["Table Grid"]
    for row_index, row in enumerate(element.rows):
        for column_index, text in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(page.size - 2 if element.borders else page.size)
            run.bold = element.header and row_index == 0
            if not element.borders and column_index == columns - 1:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph()


def _convert_to_pdf(source: Path, target: Path) -> None:
    soffice = shutil.which("soffice")
    if soffice is None:
        raise SystemExit("LibreOffice (soffice) is required to build PDF demo documents")
    with tempfile.TemporaryDirectory() as profile, tempfile.TemporaryDirectory() as out:
        subprocess.run(  # noqa: S603 - fixed arguments, local tool
            [
                soffice,
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                out,
                str(source),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
        shutil.move(Path(out) / f"{source.stem}.pdf", target)


def _scan(pdf: Path, target: Path) -> None:
    """Turns a PDF into page images as a flatbed scanner would: skew, noise, JPEG."""
    pages: list[Image.Image] = []
    document = pypdfium2.PdfDocument(str(pdf))
    for page in document:
        image = page.render(scale=SCAN_DPI / 72).to_pil().convert("L")
        image = image.rotate(SCAN_SKEW_DEGREES, resample=Image.Resampling.BICUBIC, fillcolor=250)
        noise = Image.effect_noise(image.size, SCAN_NOISE_SIGMA)
        image = Image.blend(image, noise, SCAN_NOISE_SHARE).filter(
            ImageFilter.GaussianBlur(SCAN_BLUR)
        )
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=SCAN_JPEG_QUALITY)
        pages.append(Image.open(io.BytesIO(buffer.getvalue())).convert("L"))
    document.close()
    first, *rest = pages
    first.save(target, "PDF", resolution=SCAN_DPI, save_all=True, append_images=rest)


def build(prefixes: list[str]) -> None:
    for spec in DOCUMENTS:
        if prefixes and not spec.path.startswith(tuple(prefixes)):
            continue
        target = OUTPUT / spec.path
        target.parent.mkdir(parents=True, exist_ok=True)
        if spec.output == "docx":
            _write_docx(spec, target)
        else:
            with tempfile.TemporaryDirectory() as work:
                docx_path = Path(work) / f"{target.stem}.docx"
                _write_docx(spec, docx_path)
                if spec.output == "pdf":
                    _convert_to_pdf(docx_path, target)
                else:
                    text_pdf = Path(work) / f"{target.stem}.text.pdf"
                    _convert_to_pdf(docx_path, text_pdf)
                    _scan(text_pdf, target)
        print(f"built {target.relative_to(ROOT)}")  # noqa: T201


if __name__ == "__main__":
    build(sys.argv[1:])
